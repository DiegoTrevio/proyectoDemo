"""Document parsing via Docling — extracts text, tables, images from PDFs/DOCXs."""

import logging
from dataclasses import dataclass, field

logger = logging.getLogger("agentos.tools.document_parser")


@dataclass
class ParsedDocument:
    """Structured result of document parsing."""
    text: str = ""
    tables: list[dict] = field(default_factory=list)
    images: list[dict] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)
    pages: int = 0
    error: str | None = None

    @property
    def success(self) -> bool:
        return self.error is None and bool(self.text)


def parse_document(file_path: str) -> ParsedDocument:
    """Parse a PDF or DOCX file using Docling, preserving layout.

    Args:
        file_path: Path to the document file.

    Returns:
        ParsedDocument with extracted text, tables, and images.
    """
    try:
        from docling.document_converter import DocumentConverter

        converter = DocumentConverter()
        result = converter.convert(file_path)
        doc = result.document

        # Extract full text with layout
        text = doc.export_to_markdown()

        # Extract tables
        tables = []
        for table in doc.tables:
            table_data = {
                "caption": getattr(table, "caption", ""),
                "rows": [],
            }
            df = table.export_to_dataframe()
            if df is not None:
                table_data["headers"] = list(df.columns)
                table_data["rows"] = df.values.tolist()
            tables.append(table_data)

        # Extract images (references)
        images = []
        for pic in doc.pictures:
            images.append({
                "caption": getattr(pic, "caption", ""),
                "page": getattr(pic, "prov", [{}])[0].get("page_no", 0) if hasattr(pic, "prov") and pic.prov else 0,
            })

        return ParsedDocument(
            text=text,
            tables=tables,
            images=images,
            metadata={"source": file_path},
            pages=doc.num_pages() if hasattr(doc, "num_pages") else 0,
        )

    except ImportError:
        logger.warning("Docling not installed — falling back to basic parsing")
        return _fallback_parse(file_path)
    except Exception as e:
        logger.exception("Document parsing failed: %s", file_path)
        return ParsedDocument(error=f"Parse failed: {e}")


def _fallback_parse(file_path: str) -> ParsedDocument:
    """Basic fallback parser when Docling is not available."""
    if file_path.lower().endswith(".pdf"):
        return _parse_pdf_basic(file_path)
    elif file_path.lower().endswith((".docx", ".doc")):
        return _parse_docx_basic(file_path)
    else:
        return ParsedDocument(error=f"Unsupported format: {file_path}")


def _parse_pdf_basic(file_path: str) -> ParsedDocument:
    """Basic PDF text extraction without Docling."""
    try:
        import fitz  # PyMuPDF

        doc = fitz.open(file_path)
        pages_text = []
        for page in doc:
            pages_text.append(page.get_text())
        doc.close()
        return ParsedDocument(
            text="\n\n".join(pages_text),
            pages=len(pages_text),
            metadata={"source": file_path, "parser": "pymupdf"},
        )
    except ImportError:
        return ParsedDocument(error="Neither Docling nor PyMuPDF available for PDF parsing")


def _parse_docx_basic(file_path: str) -> ParsedDocument:
    """Basic DOCX text extraction without Docling."""
    try:
        from docx import Document

        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return ParsedDocument(
            text="\n\n".join(paragraphs),
            metadata={"source": file_path, "parser": "python-docx"},
        )
    except ImportError:
        return ParsedDocument(error="Neither Docling nor python-docx available for DOCX parsing")
