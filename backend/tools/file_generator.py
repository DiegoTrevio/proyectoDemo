"""File generation tools for AgentOS — PDF, PPTX, DOCX, Excel."""

import io
import logging
from dataclasses import dataclass

logger = logging.getLogger("agentos.tools.file_generator")


@dataclass
class GeneratedFile:
    """Result of file generation."""
    name: str
    content: bytes
    mime_type: str
    format: str  # pdf, pptx, docx, xlsx

    @property
    def size_kb(self) -> float:
        return len(self.content) / 1024


def generate_pdf(
    content: str,
    title: str = "Document",
    author: str = "AgentOS",
) -> GeneratedFile:
    """Generate a PDF from text content using ReportLab."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter,
                            topMargin=inch, bottomMargin=inch,
                            leftMargin=inch, rightMargin=inch)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("CustomTitle", parent=styles["Title"], fontSize=24, spaceAfter=20)
    body_style = ParagraphStyle("CustomBody", parent=styles["Normal"], fontSize=11, leading=16)

    story = [Paragraph(title, title_style), Spacer(1, 12)]

    for paragraph in content.split("\n\n"):
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        if paragraph.startswith("# "):
            story.append(Spacer(1, 12))
            story.append(Paragraph(paragraph[2:], styles["Heading1"]))
        elif paragraph.startswith("## "):
            story.append(Spacer(1, 8))
            story.append(Paragraph(paragraph[3:], styles["Heading2"]))
        else:
            # Escape XML special chars for ReportLab
            safe = paragraph.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(safe, body_style))
            story.append(Spacer(1, 6))

    doc.build(story)
    buf.seek(0)
    return GeneratedFile(
        name=f"{title.replace(' ', '_')}.pdf",
        content=buf.read(),
        mime_type="application/pdf",
        format="pdf",
    )


def generate_pptx(
    slides_data: list[dict],
    title: str = "Presentation",
) -> GeneratedFile:
    """Generate a PPTX from slide data.

    slides_data: [{"title": str, "content": str, "image_bytes": bytes | None}, ...]
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for i, slide_data in enumerate(slides_data):
        if i == 0:
            layout = prs.slide_layouts[0]  # Title slide
        else:
            layout = prs.slide_layouts[1]  # Title + Content

        slide = prs.slides.add_slide(layout)

        # Title
        if slide.shapes.title:
            slide.shapes.title.text = slide_data.get("title", f"Slide {i + 1}")

        # Content
        content = slide_data.get("content", "")
        if content and len(slide.placeholders) > 1:
            body = slide.placeholders[1]
            body.text = content
            for paragraph in body.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(18)

        # Image
        image_bytes = slide_data.get("image_bytes")
        if image_bytes and i > 0:
            img_stream = io.BytesIO(image_bytes)
            try:
                slide.shapes.add_picture(img_stream, Inches(8), Inches(1.5),
                                         width=Inches(4.5))
            except Exception as e:
                logger.warning("Failed to add image to slide %d: %s", i + 1, e)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return GeneratedFile(
        name=f"{title.replace(' ', '_')}.pptx",
        content=buf.read(),
        mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        format="pptx",
    )


def generate_docx(
    content: str,
    title: str = "Document",
    author: str = "AgentOS",
) -> GeneratedFile:
    """Generate a DOCX from text content."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.core_properties.author = author
    doc.add_heading(title, level=0)

    for paragraph in content.split("\n\n"):
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        if paragraph.startswith("# "):
            doc.add_heading(paragraph[2:], level=1)
        elif paragraph.startswith("## "):
            doc.add_heading(paragraph[3:], level=2)
        elif paragraph.startswith("### "):
            doc.add_heading(paragraph[4:], level=3)
        elif paragraph.startswith("- ") or paragraph.startswith("* "):
            for line in paragraph.split("\n"):
                line = line.lstrip("- *").strip()
                if line:
                    doc.add_paragraph(line, style="List Bullet")
        else:
            p = doc.add_paragraph(paragraph)
            for run in p.runs:
                run.font.size = Pt(11)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return GeneratedFile(
        name=f"{title.replace(' ', '_')}.docx",
        content=buf.read(),
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        format="docx",
    )


def generate_excel(
    data: list[dict],
    sheets: dict[str, list[dict]] | None = None,
    title: str = "Data",
) -> GeneratedFile:
    """Generate an Excel file from data.

    data: list of dicts for the default sheet.
    sheets: optional {sheet_name: [row_dicts]} for multiple sheets.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()

    all_sheets = sheets or {title: data}

    for idx, (sheet_name, rows) in enumerate(all_sheets.items()):
        if idx == 0:
            ws = wb.active
            ws.title = sheet_name
        else:
            ws = wb.create_sheet(sheet_name)

        if not rows:
            continue

        # Header row
        headers = list(rows[0].keys())
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)

        # Data rows
        for row_idx, row_data in enumerate(rows, 2):
            for col, header in enumerate(headers, 1):
                ws.cell(row=row_idx, column=col, value=row_data.get(header, ""))

        # Auto-width columns
        for col_idx, header in enumerate(headers, 1):
            max_len = max(len(str(header)), *(len(str(r.get(header, ""))) for r in rows))
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 2, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return GeneratedFile(
        name=f"{title.replace(' ', '_')}.xlsx",
        content=buf.read(),
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        format="xlsx",
    )
